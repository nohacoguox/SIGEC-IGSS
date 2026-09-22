# -*- coding: utf-8 -*-
"""Utilidades python-docx para insertar el Capítulo V con estilo del informe (APA 7)."""
from __future__ import annotations

from pathlib import Path

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph

FONT = "Calibri"
SIZE = 11


def set_run(run, *, size=SIZE, bold=False, italic=False, name=FONT):
    run.font.name = name
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), "Times New Roman")
    rFonts.set(qn("w:cs"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(0, 0, 0)


def clear_para(p: Paragraph):
    for child in list(p._p):
        tag = child.tag.split("}")[-1]
        if tag in ("r", "hyperlink", "ins", "del"):
            p._p.remove(child)


def _spacing(p, *, before=0, after=0, line=None, first=None, align=None):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line == "double":
        pf.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    elif line == "single":
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if first is not None:
        pf.first_line_indent = Cm(first)
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


class Cursor:
    def __init__(self, doc, target: Paragraph):
        self.doc = doc
        self.target = target
        self.last = None

    def _new_p(self) -> Paragraph:
        el = OxmlElement("w:p")
        if self.last is None:
            self.target._p.addprevious(el)
        else:
            self.last.addnext(el)
        p = Paragraph(el, self.target._parent)
        self.last = el
        return p

    def para(self, text="", *, kind="body"):
        p = self._new_p()
        clear_para(p)
        if kind == "chapter":
            _spacing(p, before=12, after=12, first=0, align="center")
            r = p.add_run(text)
            set_run(r, bold=True)
        elif kind == "h2":
            _spacing(p, before=12, after=6, first=0, align="left")
            r = p.add_run(text)
            set_run(r, bold=True)
        elif kind == "h3":
            _spacing(p, before=10, after=6, first=0, align="left")
            r = p.add_run(text)
            set_run(r, bold=True, italic=True)
        elif kind == "caption_num":
            _spacing(p, before=12, after=0, first=0, align="left")
            r = p.add_run(text)
            set_run(r, bold=True)
        elif kind == "caption_title":
            _spacing(p, before=0, after=6, first=0, align="left")
            r = p.add_run(text)
            set_run(r, italic=True)
        elif kind == "note":
            _spacing(p, before=0, after=10, first=0, align="left")
            r1 = p.add_run("Nota. ")
            set_run(r1, italic=True)
            r2 = p.add_run(text)
            set_run(r2)
        elif kind == "center":
            _spacing(p, before=0, after=6, first=0, align="center", line="single")
        else:
            _spacing(p, before=0, after=8, first=1.27, align="justify")
            r = p.add_run(text)
            set_run(r)
        return p

    def page_break(self):
        p = self._new_p()
        clear_para(p)
        _spacing(p, first=0)
        r = p.add_run()
        r.add_break(WD_BREAK.PAGE)
        return p

    def image(self, path: Path, width_cm=15.6):
        p = self.para(kind="center")
        run = p.add_run()
        run.add_picture(str(path), width=Cm(width_cm))
        return p

    def table(self, headers, rows):
        table = self.doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        tbl = table._tbl
        parent = tbl.getparent()
        parent.remove(tbl)
        if self.last is None:
            self.target._p.addprevious(tbl)
        else:
            self.last.addnext(tbl)
        self.last = tbl

        def fill(cell, text, bold=False):
            cell.text = ""
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf = p.paragraph_format
            pf.space_before = Pt(2)
            pf.space_after = Pt(2)
            pf.first_line_indent = Cm(0)
            r = p.add_run(str(text))
            set_run(r, size=10, bold=bold)
            _set_cell_borders(cell)

        for j, h in enumerate(headers):
            fill(table.rows[0].cells[j], h, bold=True)
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                fill(table.rows[i + 1].cells[j], val, bold=False)
        _set_tbl_width(table)
        return table


def _set_cell_borders(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    for old in tcPr.findall(qn("w:tcBorders")):
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, val in (
        ("top", "single"),
        ("left", "nil"),
        ("bottom", "single"),
        ("right", "nil"),
    ):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        if val == "single":
            el.set(qn("w:sz"), "4")
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "666666")
        borders.append(el)
    tcPr.append(borders)


def _set_tbl_width(table: Table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        tbl.insert(0, tblPr)
    for old in tblPr.findall(qn("w:tblW")):
        tblPr.remove(old)
    w = OxmlElement("w:tblW")
    w.set(qn("w:w"), "5000")
    w.set(qn("w:type"), "pct")
    tblPr.append(w)


def find_para(doc, pred):
    for p in doc.paragraphs:
        if pred(p):
            return p
    raise SystemExit("Párrafo no encontrado")
