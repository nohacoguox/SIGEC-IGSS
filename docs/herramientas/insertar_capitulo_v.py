# -*- coding: utf-8 -*-
"""Inserta el Capítulo V — Diseño de la solución en el informe SIGEC-IGSS."""
from __future__ import annotations

import shutil
import sys
from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as P

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT))
from cap5_lib import Cursor, find_para  # noqa: E402

SRC = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\SEGUNDA_ENTREGA_BASE.docx")
OUT_ONEDRIVE = Path(
    r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA"
    r"\TERCERA_ENTREGA_CAPITULO_V.docx"
)
OUT_REPO = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\TERCERA_ENTREGA_CAPITULO_V.docx")

ER = ROOT.parent / "diagramas_er"
C5 = ROOT.parent / "diagramas_cap5"
CAP = C5 / "capturas"

# Numeración continua: última tabla 22, última figura de cuerpo 22; anexos 23-24 se renumeran.


def fig_file(*cands: Path) -> Path:
    for p in cands:
        if p.exists():
            return p
    raise SystemExit(f"Falta imagen: {cands[0]}")


def emit_fig(cur: Cursor, num: int, title: str, path: Path, note: str, width=15.6):
    cur.para(f"Figura {num}", kind="caption_num")
    cur.para(title, kind="caption_title")
    cur.image(path, width)
    cur.para(note, kind="note")


def emit_tbl(cur: Cursor, num: int, title: str, headers, rows, note: str):
    cur.para(f"Tabla {num}", kind="caption_num")
    cur.para(title, kind="caption_title")
    cur.table(headers, rows)
    cur.para(note, kind="note")


def write_chapter(cur: Cursor):
    cur.page_break()
    cur.para("Capítulo V – Diseño de la solución", kind="chapter")

    cur.para(
        "Este capítulo traduce la especificación del Capítulo IV en decisiones de diseño "
        "ejecutables: cómo se organizan los datos, dónde se resguardan los documentos, "
        "qué arquitectura soporta el piloto Palín–DAF Escuintla, cómo se presenta la "
        "interfaz al usuario y cómo se agrupa el software en módulos de programación. "
        "El diseño se documenta con el prototipo real SIGEC-IGSS (aplicación web React, "
        "API Node.js/Express y PostgreSQL), sin incluir scripts de creación de tablas y "
        "sin sustituir plataformas estatales de contratación. Las figuras y tablas "
        "continúan la numeración del documento (a partir de la Tabla 23 y la Figura 23)."
    )
    cur.para(
        "La exposición sigue el orden pedido en la guía de la tercera entrega: diseño de "
        "datos (modelo entidad–relación por dominios y almacenamiento de documentos), "
        "diseño de arquitectura (diagrama, detalle y especificación), diseño de la "
        "interfaz (prototipo) y programación (módulos principales y diagrama de "
        "componentes). Cada apartado cierra con la implicación práctica para el trámite "
        "de la SIAF y del expediente de compra."
    )

    # ── 5.1 ────────────────────────────────────────────────────────────────
    cur.para("5.1 Diseño de datos", kind="h2")
    cur.para(
        "El diseño de datos materializa, en un esquema relacional propio, las reglas "
        "levantadas con el actor interesado y con el personal operativo: correlativo "
        "único, estados del trámite, bitácora de rechazos y correcciones, versionado "
        "documental y alcance de consulta según unidad y permiso. La persistencia se "
        "implementa en PostgreSQL (base igss, esquema sigec_igss) mediante entidades "
        "TypeORM. El modelo físico no se presenta como script SQL, sino como modelo "
        "entidad–relación desglosado por dominios, tablas principales y reglas de "
        "integridad. El detalle atributo por atributo se concentra en el anexo de este "
        "capítulo, a fin de conservar legibilidad en el cuerpo del informe."
    )
    cur.para(
        "Se distingue el dato estructurado (filas y claves) del documento binario (PDF, "
        "Excel, imágenes). El primero vive en el motor relacional; el segundo, en el "
        "sistema de archivos, vinculado por ruta, tamaño y huella SHA-256. Esa decisión "
        "híbrida evita inflar la base con bytes y permite respaldar evidencias de forma "
        "conjunta, tal como se exige a un trámite auditable."
    )

    cur.para("5.1.1 Modelo entidad–relación", kind="h3")
    cur.para(
        "El modelo entidad–relación de SIGEC-IGSS agrupa 25 entidades persistentes y dos "
        "tablas de unión (user_roles y role_permissions). El volumen impide una única "
        "figura legible en página tamaño carta; por ello se presenta por dominios, en "
        "cinco partes secuenciales (Figuras 23 a 27). Las líneas continuas representan "
        "relaciones con integridad referencial; las discontinuas, asociaciones lógicas "
        "sin clave foránea en el ORM (por ejemplo, usuario–unidad médica por texto, "
        "ítem SIAF–catálogo por código y origen, expediente–SIAF por numero_siaf). El "
        "mismo modelo se introdujo en el Capítulo IV como apoyo a la especificación; "
        "aquí se interpreta como diseño de datos: qué se persiste, con qué cardinalidad "
        "y con qué restricciones de unicidad."
    )

    emit_fig(
        cur, 23,
        "Modelo entidad–relación de SIGEC-IGSS — Parte A: mapa de módulos",
        ER / "er_parte_a_mapa_modulos.png",
        "Organiza el esquema en cinco dominios lógicos. Las partes B–E detallan "
        "atributos, claves y cardinalidades. Elaboración propia a partir de las "
        "entidades TypeORM del prototipo (25 entidades y 2 tablas de unión).",
        16.2,
    )
    emit_tbl(
        cur, 23,
        "Dominios del modelo de datos de SIGEC-IGSS",
        ["Dominio", "Entidades principales", "Propósito de diseño"],
        [
            ["Seguridad y acceso", "user, credential, role, permission, user_roles, role_permissions",
             "Identidad, credencial con hash y RBAC granular por pantalla"],
            ["Estructura organizacional", "puesto, areas, departamentos, municipios, unidad_medica",
             "Catálogos que evitan texto libre inconsistente en SIAF y usuarios"],
            ["SIAF", "siaf_solicitudes e ítems, subproductos, autorizaciones, bitácora y adjuntos",
             "Ciclo de la solicitud de compra con correlativo, estados y evidencia"],
            ["Expedientes", "expedientes, documentos, versiones, bitácora y detalle de observaciones",
             "Carátula documental, versionado y dictamen DAF localizable en el PDF"],
            ["Catálogos y correlativos", "producto_catalogo, configuraciones y reservas de correlativo",
             "Códigos MINFIN/SIBOFA y numeración anual sin colisiones"],
        ],
        "Elaboración propia. Los nombres de tabla corresponden al esquema sigec_igss en PostgreSQL.",
    )

    cur.para(
        "Dominio de seguridad. Un usuario institucional se identifica por código de "
        "empleado, DPI, NIT y correo institucional (restricciones de unicidad). La "
        "contraseña no se guarda en la misma fila de datos personales: vive en "
        "credential, con hash y marca de clave temporal. El modelo RBAC es N:M: un "
        "usuario puede tener varios roles y un rol varios permisos, de modo que el "
        "piloto no duplica cuentas para otorgar una pantalla (por ejemplo, bandeja DAF "
        "sin administración de usuarios). El puesto y el departamento de dirección se "
        "asocian por clave foránea opcional; la unidad médica del usuario se registra "
        "como texto de catálogo, asociación lógica que el diseño reconoce y no oculta."
    )
    emit_fig(
        cur, 24,
        "Modelo entidad–relación — Parte B: seguridad, acceso y estructura organizacional",
        ER / "er_parte_b_seguridad_org.png",
        "Incluye Usuario, Credencial, Rol, Permiso, tablas de unión, Puesto, "
        "Departamento, Municipio, Unidad médica y Área. La relación Usuario–Unidad "
        "médica es lógica (texto), no clave foránea. Elaboración propia.",
        16.2,
    )

    cur.para(
        "Dominio SIAF. siaf_solicitudes es el núcleo: correlativo único, fecha, "
        "solicitante, autoridad, encargado de despacho, área, justificación, estado "
        "(borrador, pendiente, autorizado, rechazado) y bandera de aprobación de "
        "Dirección Departamental. Los ítems y subproductos se descomponen en tablas "
        "hijas 1:N para conservar orden, código de catálogo y origen (MINFIN o SIBOFA). "
        "Autorizaciones y bitácora registran quién dictaminó, con comentario, motivo "
        "estructurado y, cuando hay corrección, detalle anterior y posterior. Los "
        "adjuntos (especificaciones, cotizaciones) se modelan como metadatos; el "
        "archivo físico se describe en 5.1.2. El PDF generado deja constancia con "
        "ruta, hash y tamaño en la propia solicitud."
    )
    emit_fig(
        cur, 25,
        "Modelo entidad–relación — Parte C: módulo SIAF",
        ER / "er_parte_c_siaf.png",
        "Muestra SIAF_SOLICITUD y dependencias 1:N (ítems, subproductos, autorizaciones, "
        "bitácora y documentos adjuntos), además de vínculos con Usuario y Área. El "
        "vínculo con ProductoCatalogo es lógico por código y origen. Elaboración propia.",
        16.2,
    )

    cur.para(
        "Dominio de expedientes. El expediente agrupa la carátula (número único, tipo, "
        "título, unidad y municipio de origen, orden de compra y número SIAF asociado). "
        "Los documentos hijos soportan tipos institucionales (SIAF, factura, orden de "
        "compra, contrato). El versionado (expediente_documento_versiones) conserva cada "
        "reemplazo con es_actual, de manera que un rechazo no borre la evidencia de lo "
        "revisado. La bitácora y su detalle permiten anotar motivo por documento, página "
        "y posición en el visor PDF: diseño alineado con la auditabilidad exigida en el "
        "levantamiento de requerimientos, no con un simple repositorio de archivos."
    )
    emit_fig(
        cur, 26,
        "Modelo entidad–relación — Parte D: módulo de expedientes",
        ER / "er_parte_d_expedientes.png",
        "Incluye Expediente, documentos, versiones, bitácora y detalle de observaciones "
        "documentales. El campo numero_siaf vincula lógicamente el expediente con una "
        "solicitud SIAF sin clave foránea en el ORM. Elaboración propia.",
        16.2,
    )

    cur.para(
        "Dominio de catálogos y correlativos. producto_catalogo guarda códigos y "
        "descripciones por origen, con la fila original en JSONB para no perder columnas "
        "del Excel oficial. La configuración de mapeo (columna de código y columnas de "
        "descripción) se separa por origen. Los correlativos SIAF y de expedientes tienen "
        "configuración de siguiente número, dígitos y año calendario; las reservas "
        "temporales evitan que dos usuarios tomen el mismo correlativo mientras llenan el "
        "formulario. Si la reserva expira o se cancela, el número se libera. Ese diseño "
        "responde a un riesgo operativo real en consultorio, no a un adorno de modelo."
    )
    emit_fig(
        cur, 27,
        "Modelo entidad–relación — Parte E: catálogo de productos y correlativos",
        ER / "er_parte_e_catalogos_correlativos.png",
        "Cubre ProductoCatalogo, ProductoCatalogoConfig, configuración y reserva de "
        "correlativos SIAF, y configuración de correlativos de expedientes. Elaboración propia.",
        16.2,
    )

    emit_tbl(
        cur, 24,
        "Tablas principales del modelo físico y su responsabilidad",
        ["Tabla", "Clave de negocio", "Responsabilidad"],
        [
            ["user / credential", "codigoEmpleado", "Persona y secreto de acceso separados"],
            ["role / permission", "name", "Perfiles y facultades de pantalla"],
            ["siaf_solicitudes", "correlativo", "Cabecera del trámite SIAF y PDF"],
            ["siaf_items / siaf_subproductos", "siaf_id + orden", "Líneas de bien y subproducto"],
            ["siaf_bitacora / siaf_autorizaciones", "siaf_id + fecha", "Dictamen y traza"],
            ["expedientes", "numeroExpediente", "Carátula del expediente de compra"],
            ["expediente_documentos / versiones", "expediente_id + n.º versión", "Archivo vigente y respaldos"],
            ["producto_catalogo", "origen + codigo", "Catálogo MINFIN / SIBOFA / subproductos"],
            ["siaf_correlativo_reservas", "numero + estado", "Anti-colisión de correlativos"],
        ],
        "Elaboración propia. No se incluyen scripts DDL; la descripción detallada de atributos se presenta en la Tabla 38 (anexo).",
    )
    emit_tbl(
        cur, 25,
        "Relaciones y cardinalidades relevantes del diseño",
        ["Origen", "Destino", "Cardinalidad", "Implementación"],
        [
            ["user", "role", "N:M", "Tabla de unión user_roles"],
            ["role", "permission", "N:M", "Tabla de unión role_permissions"],
            ["credential", "user", "1:1", "FK userId con borrado en cascada"],
            ["siaf_solicitudes", "user", "N:1 (×3)", "Solicitante, autoridad y encargado"],
            ["siaf_* hijas", "siaf_solicitudes", "N:1", "FK siaf_id; cascada en ítems, bitácora y adjuntos"],
            ["municipios", "departamentos", "N:1", "FK departamento_id"],
            ["expediente_documentos", "expedientes", "N:1", "FK expediente_id en cascada"],
            ["expediente_documento_versiones", "expediente_documentos", "N:1", "Historial de reemplazos"],
            ["siaf_items", "producto_catalogo", "N:1 lógica", "codigo + catalogo_origen, sin FK ORM"],
            ["expedientes", "siaf_solicitudes", "N:1 lógica", "numero_siaf, sin FK ORM"],
        ],
        "Elaboración propia. Las asociaciones lógicas se documentan para no simular integridad referencial donde el prototipo opera por coincidencia de negocio.",
    )
    emit_tbl(
        cur, 26,
        "Reglas de integridad adoptadas en el diseño de datos",
        ["Regla", "Dónde aplica", "Efecto en el trámite"],
        [
            ["Unicidad de correlativo / número de expediente", "siaf_solicitudes, expedientes", "Identificador visible y no reutilizable en el año"],
            ["Unicidad de DPI, NIT, correo y código de empleado", "user", "Evita dobles padrones"],
            ["Unicidad origen+código", "producto_catalogo", "Impide códigos duplicados por catálogo"],
            ["Cascada de hijas SIAF y expediente", "ítems, adjuntos, versiones, bitácoras", "No quedan huérfanos al eliminar el padre en pruebas controladas"],
            ["Hash SHA-256 de archivo", "PDF SIAF, adjuntos, documentos", "Detecta corrupción o sustitución en disco"],
            ["Estados controlados", "siaf_solicitudes.estado, expedientes.estado", "La analítica y las bandejas filtran con valores estables"],
            ["Reserva temporal de correlativo", "siaf_correlativo_reservas", "Dos formularios abiertos no consumen el mismo número"],
        ],
        "Elaboración propia a partir del modelo TypeORM y de los servicios de correlativo y almacenamiento.",
    )
    cur.para(
        "En conclusión, el modelo entidad–relación no es un dibujo ornamental: es el "
        "contrato de memoria del trámite Consultorio–DAF. Desglosarlo por dominios "
        "permite auditar seguridad, SIAF, expedientes y correlativos por separado, y "
        "deja explícitas las asociaciones lógicas que el prototipo usa sin fingir una "
        "clave foránea inexistente."
    )

    cur.para("5.1.2 Almacenamiento de documentos", kind="h3")
    cur.para(
        "SIGEC-IGSS adopta almacenamiento híbrido: PostgreSQL guarda metadatos y el "
        "disco del servidor guarda bytes. El servicio FileStorageService concentra la "
        "escritura, lectura, borrado y cálculo de hash. Los correlativos con formato "
        "número/año se sanitizan para no interpretar la barra como subdirectorio. Esta "
        "separación permite consultas rápidas (estado, fechas, responsables) sin abrir "
        "PDF, y conserva la evidencia imprimible que el IGSS sigue necesitando."
    )
    emit_fig(
        cur, 28,
        "Esquema de almacenamiento híbrido de documentos en SIGEC-IGSS",
        C5 / "almacenamiento_hibrido.png",
        "Los metadatos (ruta, hash, tamaño, MIME, usuario y fecha) permanecen en "
        "PostgreSQL; los archivos se organizan en uploads/siaf, uploads/siaf-adjuntos "
        "y uploads/expedientes. Elaboración propia.",
    )
    emit_tbl(
        cur, 27,
        "Estructura de directorios del almacén documental",
        ["Ruta relativa", "Contenido", "Clave de organización"],
        [
            ["uploads/siaf/AAAA/MM/", "PDF institucional de la solicitud SIAF", "Año y mes de generación"],
            ["uploads/siaf-adjuntos/AAAA/MM/", "Especificaciones, cotizaciones u otros soportes", "Identificador SIAF + marca de tiempo"],
            ["uploads/expedientes/<número>/", "Documentos y versiones del expediente", "Número sanitizado del expediente"],
        ],
        "Elaboración propia. Las rutas relativas se persisten en pdfPath o rutaArchivo; el respaldo debe copiar este árbol junto con el volcado de la base.",
    )
    emit_tbl(
        cur, 28,
        "Políticas de archivo aplicadas en el prototipo",
        ["Política", "Mecanismo", "Límite o nota del piloto"],
        [
            ["Integridad", "SHA-256 al persistir", "Permite verificar el fichero restaurado"],
            ["Unicidad de nombre", "Marca de tiempo en adjuntos y versiones", "Evita sobrescribir silenciosamente"],
            ["Versionado", "expediente_documento_versiones.es_actual", "DAF revisa la versión señalada, no un archivo anónimo"],
            ["Autorización de descarga", "API con JWT y permiso", "El disco no se publica como carpeta web abierta"],
            ["Tamaño de carga", "Límite de multer / catálogo Excel", "Debe alinearse a política institucional"],
            ["Restauración", "Base + uploads en el mismo punto de corte", "Restaurar solo uno deja metadatos huérfanos"],
        ],
        "Elaboración propia. El diseño no constituye un gestor documental certificado; sí cubre integridad, traza y versionado exigidos por el piloto.",
    )
    cur.para(
        "En conclusión, el almacenamiento de documentos convierte el PDF y los adjuntos "
        "en evidencia ligada al expediente o a la SIAF, no en archivos sueltos de "
        "escritorio. La huella digital, la ruta relativa y la bitácora forman un todo "
        "que debe respaldarse junto; de lo contrario se rompe el valor probatorio que "
        "justifica digitalizar el trámite."
    )

    # ── 5.2 ────────────────────────────────────────────────────────────────
    cur.para("5.2 Diseño de la arquitectura", kind="h2")
    cur.para(
        "La arquitectura del SIGEC-IGSS se diseña como sistema cliente–servidor en tres "
        "capas (presentación, aplicación y persistencia), desplegado en un servidor "
        "Debian 12 de la unidad piloto. El Capítulo III describió el marco técnico; "
        "este apartado lo convierte en vistas de diseño: contenedores, despliegue, "
        "preocupaciones de calidad y especificación de decisiones. Se consideran los "
        "aspectos que una descripción arquitectónica debe hacer visibles: interesados "
        "y preocupaciones, vistas (lógica, de proceso, de desarrollo y física), "
        "atributos de calidad, decisiones, interfaces y límites del piloto, en línea "
        "con la práctica de documentar arquitecturas de software y con el modelo de "
        "vistas 4+1."
    )

    cur.para("5.2.1 Diagrama de arquitectura de la solución", kind="h3")
    cur.para(
        "La Figura 29 muestra la vista de contenedores. Los actores (colaborador, "
        "analista DAF y administrador) usan un navegador. La SPA React no persiste "
        "reglas de negocio: consume la API REST en JSON con token JWT. Nginx es el "
        "punto único de entrada: entrega los estáticos compilados y encamina /api al "
        "proceso Node supervisado por PM2. La API orquesta módulos de dominio y "
        "servicios (almacenamiento, PDF, correo, Excel, apoyo lingüístico). PostgreSQL "
        "y el directorio uploads/ cierran la persistencia."
    )
    emit_fig(
        cur, 29,
        "Arquitectura de la solución SIGEC-IGSS (vista de contenedores)",
        C5 / "arquitectura_solucion.png",
        "Separación entre SPA, API REST, motor relacional y almacén de archivos. "
        "Elaboración propia a partir del prototipo implementado.",
        16.0,
    )
    cur.para(
        "La Figura 30 concreta el despliegue del piloto: IP fija 10.4.201.74, HTTP en "
        "puerto 80, API solo en localhost:3001, UFW con SSH y HTTP, y secretos fuera de "
        "Git. El mismo repositorio se desarrolla en Windows y se publica en Debian; la "
        "URL del API queda congelada en el build del frontend, de ahí la exigencia de "
        "dirección estable."
    )
    emit_fig(
        cur, 30,
        "Arquitectura de despliegue del piloto SIGEC-IGSS (Palín–DAF Escuintla)",
        C5 / "arquitectura_despliegue.png",
        "Servidor Debian 12 (Dell OptiPlex Micro) con Nginx, PM2, PostgreSQL y disco "
        "de uploads. Elaboración propia según la guía de despliegue del proyecto.",
        16.0,
    )

    cur.para("5.2.2 Detalle de la arquitectura", kind="h3")
    cur.para(
        "Interesados y preocupaciones. El colaborador del consultorio necesita captura "
        "ágil y corrección tras un rechazo. El analista DAF necesita bandeja, motivo "
        "estructurado y evidencia de lo revisado. El administrador necesita padrones y "
        "correlativos. Sistemas/IT necesita un único punto de publicación, respaldos y "
        "secretos fuera del código. El tribunal académico necesita trazabilidad entre "
        "requerimiento, diseño y prototipo. Esas preocupaciones gobiernan las vistas "
        "siguientes; no se diseña “para el diagrama”, se diseña para esas tensiones."
    )
    emit_fig(
        cur, 31,
        "Vistas de arquitectura del SIGEC-IGSS según el modelo 4+1",
        C5 / "vistas_arquitectura.png",
        "Vistas lógica, de proceso, de desarrollo y física, más escenarios de casos de "
        "uso (UC-01 a UC-11) y preocupaciones de calidad. Elaboración propia con base "
        "en el modelo de vistas 4+1.",
        16.0,
    )
    emit_tbl(
        cur, 29,
        "Vistas arquitectónicas y preocupaciones cubiertas",
        ["Vista", "Qué muestra", "Preocupación que atiende"],
        [
            ["Lógica", "Módulos de negocio y entidades", "Adecuación al trámite SIAF/expediente"],
            ["Proceso", "Login, reserva, dictamen, versionado", "Trazabilidad y concurrencia de correlativos"],
            ["Desarrollo", "frontend/, backend/src/modules, entity, services", "Mantenibilidad y traspaso a sistemas"],
            ["Física", "Debian, Nginx, PM2, PostgreSQL, uploads", "Disponibilidad del piloto en LAN IGSS"],
            ["Escenarios", "Casos de uso UC-01 a UC-11", "Validación de que la arquitectura ejecuta la especificación"],
        ],
        "Elaboración propia. La tabla resume el detalle de la Figura 31.",
    )
    cur.para(
        "Atributos de calidad. La seguridad se apoya en JWT, hash de credenciales, RBAC "
        "en servidor y no exposición directa del disco. La usabilidad se apoya en SPA, "
        "Material UI y menús filtrados por permiso. La mantenibilidad se apoya en "
        "TypeScript, módulos de API y un repositorio único. La integridad documental se "
        "apoya en hash y versiones. El rendimiento del piloto se sostiene con agregados "
        "calculados en servidor (analítica) y no con arrastrar el histórico al "
        "navegador. La continuidad exige RTO/RPO institucionales; el diseño solo ofrece "
        "el medio técnico (volcado + copia de uploads)."
    )
    emit_fig(
        cur, 32,
        "Flujo de una petición autenticada a través de las capas de SIGEC-IGSS",
        C5 / "flujo_capas_peticion.png",
        "El navegador nunca es la autoridad: Nginx, Express (verifyToken y permisos), "
        "servicios de dominio y persistencia ejecutan el caso de uso. Elaboración propia.",
        16.0,
    )
    cur.para(
        "Patrones. Cliente–servidor, capas, REST, SPA, RBAC, almacenamiento híbrido y "
        "proxy inverso. No se introduce microservicios ni contenedores de orquestación: "
        "la escala del piloto (una unidad y su DAF) no justifica esa complejidad. Las "
        "deudas reconocidas —token en almacenamiento del navegador, CORS amplio en "
        "desarrollo, pruebas automatizadas aún mínimas en backend, HTTPS pendiente de "
        "certificado institucional— se declaran para no vender madurez inexistente."
    )
    cur.para(
        "En conclusión, el detalle arquitectónico exhibe capas, vistas, calidades y "
        "límites. Sirve tanto al tribunal (juzgar coherencia) como a sistemas del IGSS "
        "(operar y respaldar) sin confundir el piloto con una certificación ISO 27001."
    )

    cur.para("5.2.3 Especificación de la arquitectura", kind="h3")
    cur.para(
        "La especificación fija tecnologías, interfaces y decisiones que no deben "
        "quedar solo en prosa. La Tabla 30 declara el stack por capa; la Tabla 31, las "
        "decisiones con su justificación y consecuencia; la Tabla 32, el contrato de "
        "módulos de la API que la interfaz consume."
    )
    emit_tbl(
        cur, 30,
        "Especificación tecnológica por capa",
        ["Capa", "Tecnología", "Responsabilidad especificada"],
        [
            ["Presentación", "React 18, TypeScript, MUI, React Router, Axios, Recharts",
             "SPA, rutas protegidas, temas institucionales, gráficos de gestión"],
            ["Aplicación", "Node.js, Express, TypeORM, JWT, multer, pdfkit, xlsx",
             "Reglas, autorización, archivos, PDF, importación y agregados"],
            ["Datos", "PostgreSQL (igss / sigec_igss)", "ACID, unicidad, JSONB de catálogo"],
            ["Archivos", "Sistema de archivos local (uploads/)", "PDF, adjuntos y versiones"],
            ["Publicación", "Nginx + PM2 + UFW en Debian 12", "Entrada única, proceso vivo, puertos acotados"],
            ["Identidad", "JWT (RFC 7519) + RBAC de pantallas", "Sesión y facultades verificadas en servidor"],
            ["Integraciones", "Correo SMTP, LanguageTool, Excel MINFIN/SIBOFA", "Recuperación de acceso, texto y catálogo"],
        ],
        "Elaboración propia. Versiones exactas constan en package.json del repositorio versionado.",
    )
    emit_tbl(
        cur, 31,
        "Decisiones de arquitectura del piloto",
        ["Decisión", "Alternativa descartada", "Justificación"],
        [
            ["Monolito modular (un repositorio, módulos de API)", "Microservicios",
             "Un equipo y un piloto; menor costo operativo"],
            ["SPA + API REST", "Renderizado exclusivo en servidor",
             "Formularios largos SIAF y paneles por rol sin recargar todo el sitio"],
            ["PostgreSQL dedicado", "Access / hojas de cálculo como sistema",
             "Integridad, concurrencia y consultas de analítica"],
            ["Archivos en disco + metadatos en BD", "BLOB en la base o solo carpeta compartida",
             "Consultas ágiles e integridad con hash"],
            ["Nginx como fachada", "Exponer Node a la LAN",
             "Un puerto, estáticos y futuro HTTPS en un solo punto"],
            ["RBAC en servidor", "Ocultar botones solo en el cliente",
             "Cualquier cliente podría forjar peticiones"],
            ["Correlativo con reserva", "MAX(número)+1 sin bloqueo",
             "Evita colisiones con formularios simultáneos"],
        ],
        "Elaboración propia. Cada decisión se puede rastrear en código y en la guía de despliegue.",
    )
    emit_tbl(
        cur, 32,
        "Especificación de interfaces de la API por módulo",
        ["Módulo", "Prefijo", "Capacidad de interfaz"],
        [
            ["auth", "/api/auth", "Login, sesión, cambio y recuperación de contraseña"],
            ["rbac", "/api/roles, /permissions", "Roles, permisos y catálogo de pantallas"],
            ["usuarios", "/api/users", "CRUD, roles asignados, consultas por unidad"],
            ["catalogos", "/api/puestos, areas, unidades, geografía", "Maestros organizacionales"],
            ["siaf", "/api/siaf", "Ciclo completo, adjuntos, bitácora, dictamen DAF"],
            ["expedientes", "/api/expedientes", "Carátula, documentos, versiones, dictamen"],
            ["correlativos", "/api/correlativos", "Configuración, reserva y consumo"],
            ["catalogoProductos", "/api/catalogo-productos", "Consulta e importación Excel"],
            ["estadisticas", "/api/estadisticas", "Cierre, tiempos, motivos y filtros de alcance"],
        ],
        "Elaboración propia. El contrato es JSON sobre HTTP; la autorización se revalida en cada recurso sensible.",
    )
    cur.para(
        "Restricciones. El piloto opera en red institucional, con HTTP hasta que exista "
        "certificado; no integra Guatecompras ni el core financiero del IGSS; no "
        "garantiza alta disponibilidad multinodo. La especificación deja esas puertas "
        "cerradas a propósito: ampliarlas sería otro proyecto. En conclusión, la "
        "arquitectura queda especificada en capas, decisiones e interfaces, de modo que "
        "un área de sistemas pueda operarla y un tribunal pueda contrastarla con el "
        "repositorio."
    )

    # ── 5.3 ────────────────────────────────────────────────────────────────
    cur.para("5.3 Diseño de la interfaz", kind="h2")
    cur.para(
        "El diseño de interfaz del SIGEC-IGSS no es un mockup estático ajeno al código: "
        "es el prototipo operativo en navegador, con identidad visual del IGSS (azul, "
        "verde, gris y blanco), tipografía de interfaz y menús que se encienden o se "
        "apagan según permisos. El objetivo es reducir carga cognitiva en jornadas "
        "administrativas largas y hacer visible el estado del trámite (pendiente, "
        "rechazado, autorizado) sin adivinar en un libro de actas. Las capturas de esta "
        "sección corresponden al prototipo en ejecución."
    )
    emit_tbl(
        cur, 33,
        "Principios de diseño de interfaz aplicados en el prototipo",
        ["Principio", "Cómo se materializa"],
        [
            ["Identidad institucional", "Paleta IGSS, logo y textos en español de Guatemala"],
            ["Autorización visible y real", "La UI oculta módulos sin permiso; la API vuelve a comprobar"],
            ["Dos paneles, un sistema", "Administrador y colaborador; cruce cuando el rol lo permite"],
            ["Formulario SIAF como documento", "Secciones del SIAF-A-01 (unidad, ítems, autoridad, justificación, PDF)"],
            ["Evidencia a la vista", "Estados, bandejas DAF, visor PDF y analítica de cierre"],
            ["Recuperación de acceso", "¿Olvidaste tu contraseña? y cambio obligatorio si la clave es temporal"],
            ["Adaptación al puesto de trabajo", "SPA en navegador; no se instala cliente pesado en cada PC"],
        ],
        "Elaboración propia a partir de theme/institutionalColors, App.tsx y el catálogo APP_SCREENS.",
    )
    emit_fig(
        cur, 33,
        "Mapa de navegación del prototipo de interfaz SIGEC-IGSS",
        C5 / "mapa_navegacion.png",
        "Desde el login, el usuario accede al panel administrativo, al panel de "
        "colaborador o a ambos, según roles. Elaboración propia.",
        16.0,
    )
    emit_tbl(
        cur, 34,
        "Pantallas del prototipo y actor principal",
        ["Pantalla", "Ruta o ubicación", "Actor", "Caso de uso"],
        [
            ["Inicio de sesión", "/login", "Todos", "UC-01, UC-11"],
            ["Panel de administración", "/admin-dashboard", "Administrador", "UC-02, UC-03"],
            ["Gestión de usuarios y roles", "Gestiones del panel admin", "Administrador", "UC-02"],
            ["Panel del colaborador", "/colaborador-dashboard", "Colaborador / DAF", "Acceso a UC-04–UC-10"],
            ["Libro SIAF", "/siaf-book", "Colaborador", "UC-04, UC-06"],
            ["Crear / corregir SIAF", "/siaf-book/crear", "Colaborador", "UC-04, UC-06"],
            ["Expedientes de compras", "/expedientes", "Colaborador", "UC-07"],
            ["Bandeja de revisiones DAF", "Vista interna del panel", "Analista DAF", "UC-05, UC-08"],
            ["Analítica", "Estadísticas del panel", "DAF / jefatura con permiso", "UC-10"],
            ["Catálogo de productos", "/actualizar-codigos-productos", "Administrador autorizado", "UC-09"],
        ],
        "Elaboración propia. Las Figuras 34 a 43 ilustran estas pantallas en el prototipo en ejecución.",
    )

    emit_fig(
        cur, 34,
        "Prototipo de interfaz — inicio de sesión",
        fig_file(CAP / "01_inicio_sesion.png", C5 / "ui_login.png"),
        "Pantalla pública. Solicita código de empleado y contraseña, ofrece recuperación "
        "de acceso y presenta la identidad visual institucional. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 35,
        "Prototipo de interfaz — panel de administración",
        fig_file(CAP / "02_panel_administracion.png", C5 / "ui_admin.png"),
        "Indicadores de usuarios, roles e informes, menú lateral de gestiones y salto al "
        "panel de colaborador. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 36,
        "Prototipo de interfaz — gestión de usuarios",
        fig_file(CAP / "03_gestion_usuarios.png", C5 / "ui_admin.png"),
        "Padrón de usuarios institucionales: alta, edición y asignación de roles y unidad. "
        "Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 37,
        "Prototipo de interfaz — panel de control del colaborador",
        fig_file(CAP / "09_panel_colaborador.png", C5 / "ui_colaborador.png"),
        "Punto de entrada operativo hacia Libro SIAF, expedientes, bandeja DAF, catálogo "
        "y analítica, filtrado por permisos. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 38,
        "Prototipo de interfaz — listado de solicitudes SIAF",
        fig_file(CAP / "10_listado_siaf.png", C5 / "ui_siaf.png"),
        "Inventario de solicitudes con estado y acciones de seguimiento. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 39,
        "Prototipo de interfaz — creación de solicitud SIAF (formato SIAF-A-01)",
        fig_file(CAP / "11_crear_siaf.png", C5 / "ui_siaf.png"),
        "Formulario de captura alineado al documento institucional: unidad ejecutora, "
        "ítems, subproductos, solicitante, autoridad, adjuntos, justificación y generación "
        "de PDF. Captura del prototipo SIGEC-IGSS.",
        14.5,
    )
    emit_fig(
        cur, 40,
        "Prototipo de interfaz — expedientes de compras",
        fig_file(CAP / "12_expedientes_compras.png", C5 / "ui_expedientes.png"),
        "Gestión de carátulas, documentos, versiones y seguimiento del ciclo ante DAF. "
        "Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 41,
        "Prototipo de interfaz — bandeja de revisiones DAF",
        fig_file(CAP / "14_bandeja_revisiones_daf.png", C5 / "ui_daf.png"),
        "Cola de trabajo del analista departamental para dictaminar SIAF y expedientes. "
        "Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 42,
        "Prototipo de interfaz — análisis SIAF (cierre, trazabilidad y motivos)",
        fig_file(CAP / "16_analisis_siaf_generales.png", C5 / "ui_analitica.png"),
        "Tablero de cierre mensual y evidencias cuantitativas del piloto. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    emit_fig(
        cur, 43,
        "Prototipo de interfaz — actualización de códigos y productos",
        fig_file(CAP / "13_actualizar_codigos_productos.png", C5 / "ui_catalogo.png"),
        "Importación controlada del catálogo desde Excel oficial. Captura del prototipo SIGEC-IGSS.",
        15.8,
    )
    cur.para(
        "El prototipo contempla, además, gestión de roles, áreas, puestos, unidades "
        "médicas y correlativos en el panel administrativo, analítica por caso y por "
        "motivos de rechazo, y cambio de contraseña tras una clave temporal. Esas "
        "pantallas siguen el mismo lenguaje visual y no se reproducen todas aquí para "
        "evitar reiteración; el mapa de la Figura 33 y la Tabla 34 las sitúan en la "
        "navegación. En dispositivos angostos, la SPA reorganiza tarjetas y menús "
        "(Material UI); el diseño de referencia del piloto es el puesto de trabajo con "
        "navegador de escritorio en la LAN institucional."
    )
    cur.para(
        "En conclusión, la interfaz hace visible el flujo Consultorio–DAF con la "
        "identidad del IGSS y con un prototipo ejecutable, no con láminas desconectadas "
        "del software. Cada pantalla de las Figuras 34 a 43 corresponde a uno o más "
        "casos de uso del Capítulo IV."
    )

    # ── 5.4 ────────────────────────────────────────────────────────────────
    cur.para("5.4 Programación", kind="h2")
    cur.para(
        "La programación del SIGEC-IGSS se organiza como monolito modular: una SPA y "
        "una API en el mismo repositorio, con módulos de servidor por dominio y "
        "componentes de interfaz reutilizables. Esta sección identifica los módulos "
        "principales y presenta el diagrama de componentes que relaciona páginas, "
        "servicios y persistencia. No se listan miles de líneas de código; se documenta "
        "la descomposición que permite mantener, desplegar y explicar el prototipo."
    )
    emit_fig(
        cur, 44,
        "Diagrama de componentes de programación de SIGEC-IGSS",
        C5 / "diagrama_componentes.png",
        "Componentes de la SPA (pages, components, context, api) y del servidor "
        "(modules, entity, services, middleware, migraciones). Elaboración propia a "
        "partir de la estructura de código del prototipo.",
        16.0,
    )
    emit_tbl(
        cur, 35,
        "Módulos principales de programación en el backend",
        ["Módulo", "Ubicación", "Función"],
        [
            ["auth", "backend/src/modules/auth", "Identidad, JWT, recuperación y cambio de clave"],
            ["rbac", "backend/src/modules/rbac", "Roles, permisos y sincronización de pantallas"],
            ["usuarios", "backend/src/modules/usuarios", "Padrón y asignación organizacional"],
            ["catalogos", "backend/src/modules/catalogos", "Puestos, áreas, unidades y geografía"],
            ["siaf", "backend/src/modules/siaf", "Ciclo de vida de la solicitud de compra"],
            ["expedientes", "backend/src/modules/expedientes", "Carátula, documentos, versiones y dictamen"],
            ["correlativos", "backend/src/modules/correlativos", "Secuencia anual y reservas"],
            ["catalogoProductos", "backend/src/modules/catalogoProductos", "Importación y consulta de códigos"],
            ["estadisticas", "backend/src/modules/estadisticas", "Indicadores de SIAF, expedientes y DAF"],
            ["FileStorage / PdfGenerator", "backend/src/services", "Disco, hash y constancia PDF"],
        ],
        "Elaboración propia. index.ts compone los enrutadores; data-source.ts declara las entidades.",
    )
    emit_tbl(
        cur, 36,
        "Módulos principales de programación en el frontend",
        ["Módulo / componente", "Ubicación", "Función"],
        [
            ["LoginPage / ChangePasswordPage", "pages/", "Acceso y credencial temporal"],
            ["AdminDashboard", "pages/", "Panel, gestiones maestras y estadísticas de control"],
            ["CollaboratorDashboard", "pages/", "Operación, bandeja DAF y analítica"],
            ["SiafManagement / SiafBook", "pages/ y components/", "Listado, captura, corrección y PDF"],
            ["ExpedientesPage", "pages/", "Expediente de compras y documentos"],
            ["BandejaRevisionesDaf", "components/", "Dictamen departamental"],
            ["AnaliticaDaf / AnaliticaExpedientes", "components/", "Cierre, tiempos y motivos"],
            ["UserManagement / RoleManagement", "components/ y pages/", "Padrones de seguridad"],
            ["PrivateRoute / usePermissions", "components/ y hooks/", "Protección de rutas y menús"],
            ["api.ts + contextos", "src/", "Cliente HTTP, tema y notificaciones"],
        ],
        "Elaboración propia a partir de App.tsx y del árbol frontend/src.",
    )
    emit_tbl(
        cur, 37,
        "Trazabilidad entre módulos de programación y requerimientos",
        ["Módulos", "Requerimientos / casos de uso"],
        [
            ["auth + LoginPage", "RF de autenticación; UC-01 y UC-11"],
            ["rbac + usuarios + gestiones admin", "RF de administración; UC-02 y UC-03"],
            ["siaf + SiafBook + correlativos", "RF de SIAF; UC-04, UC-05 y UC-06"],
            ["expedientes + visor PDF", "RF documental; UC-07 y UC-08"],
            ["catalogoProductos", "RF de catálogo; UC-09"],
            ["estadisticas + analítica UI", "RF de indicadores; UC-10; reportes RO/RC del Capítulo IV"],
            ["FileStorage + entidades de archivo", "RNF de integridad y trazabilidad documental"],
        ],
        "Elaboración propia. La matriz cierra el circuito entre el Capítulo IV y el software versionado.",
    )
    cur.para(
        "Criterios de implementación. TypeScript reduce errores de contrato entre UI y "
        "API. Las migraciones y ensureSchema evolucionan el esquema sin depender de "
        "anotaciones sueltas. Los seeders permiten levantar un ambiente de piloto con "
        "roles y catálogos. ESLint y Jest existen en el frontend; el backend del piloto "
        "aún privilegia verificación manual y compilación estática, deuda declarada en "
        "el Capítulo III. En conclusión, la programación se presenta como módulos "
        "nombrables, componentes acoplados por JSON y una traza explícita hacia los "
        "requerimientos: eso es lo que un área de sistemas puede mantener y lo que un "
        "tribunal puede evaluar sin perderse en el árbol de archivos."
    )

    cur.para("Conclusión del capítulo", kind="h2")
    cur.para(
        "El Capítulo V deja diseñado el SIGEC-IGSS en cuatro frentes exigidos por la "
        "entrega: (1) un modelo de datos por dominios, con integridad, correlativos y "
        "asociaciones lógicas declaradas; (2) un almacén documental híbrido con hash y "
        "versionado; (3) una arquitectura cliente–servidor en tres capas, vista de "
        "despliegue Palín–Escuintla y especificación de decisiones e interfaces; (4) un "
        "prototipo de interfaz real, con navegación por roles, y módulos de "
        "programación trazables a los casos de uso. El diseño permanece acotado al "
        "piloto: no reemplaza Guatecompras ni certifica el gobierno de datos del "
        "instituto, pero sí convierte el trámite Consultorio–DAF en un sistema "
        "inteligible, auditable y desplegable."
    )


def annex_dictionary(cur: Cursor):
    cur.para("Anexo del Capítulo V — Diccionario de tablas del modelo de datos", kind="h2")
    cur.para(
        "La Tabla 38 describe cada tabla persistente del prototipo. No se incluyen "
        "sentencias SQL. Los tipos se expresan en lenguaje de diseño (texto, entero, "
        "fecha, booleano, JSON) para facilitar la lectura institucional."
    )
    emit_tbl(
        cur, 38,
        "Diccionario de tablas del esquema sigec_igss",
        ["Tabla", "Atributos esenciales", "Observación de diseño"],
        [
            ["user", "id, nombres, apellidos, dpi, nit, telefono, correoInstitucional, codigoEmpleado, renglon, unidadMedica, puestoId, departamento_direccion_id",
             "Unicidad en DPI, NIT, correo y código"],
            ["credential", "id, codigoEmpleado, password (hash), isTempPassword, userId",
             "1:1 con user; la clave nunca va en texto claro"],
            ["role", "id, name", "Nombre único de perfil"],
            ["permission", "id, name, description", "Facultad de pantalla o acción"],
            ["user_roles", "userId, roleId", "Unión N:M"],
            ["role_permissions", "roleId, permissionId", "Unión N:M"],
            ["puesto", "id, nombre, activo", "Catálogo de puestos"],
            ["areas", "id, nombre, descripcion, activo, timestamps", "Área solicitante SIAF"],
            ["departamentos", "id, nombre", "Geografía / DAF"],
            ["municipios", "id, nombre, departamento_id", "N:1 con departamento"],
            ["unidad_medica", "id, nombre, departamento, municipio_id, telefonos", "Catálogo de unidades"],
            ["siaf_solicitudes", "id, correlativo, fecha, FKs de usuarios y área, justificación, estado, aprobado_dd, pdfPath, pdfHash, pdfSize",
             "Núcleo del trámite SIAF"],
            ["siaf_items", "id, siaf_id, codigo, catalogo_origen, descripcion, cantidad, orden",
             "Línea de bien o servicio"],
            ["siaf_subproductos", "id, siaf_id, codigo, cantidad, orden", "Desglose por subproducto"],
            ["siaf_autorizaciones", "id, siaf_id, usuario_autorizador_id, accion, comentario, motivo_rechazo, fecha",
             "Dictamen de unidad o DAF"],
            ["siaf_bitacora", "id, siaf_id, tipo, usuario_id, comentario, detalle_antes/después, fecha",
             "Traza de rechazo, corrección y autorización"],
            ["siaf_documentos_adjuntos", "id, siaf_id, nombreOriginal, rutaArchivo, mimeType, tamanioBytes, hashArchivo, fecha_subida",
             "Metadatos del adjunto en disco"],
            ["siaf_correlativo_config", "siguiente_numero, numero_inicio, digitos, minutos_reserva, anio_actual",
             "Una fila de configuración"],
            ["siaf_correlativo_reservas", "numero, correlativo, usuario_id, estado, expira_en",
             "reservado / consumido / liberado"],
            ["expedientes", "id, numeroExpediente, usuario_id, tipo, titulo, numero_siaf, numero_orden_compra, estado, unidad_origen, municipio_origen",
             "Carátula; vínculo lógico con SIAF"],
            ["expediente_documentos", "id, expediente_id, tipoDocumento, nombreArchivo, ruta, mime, hash, subido_por",
             "Documento vigente"],
            ["expediente_documento_versiones", "id, expediente_documento_id, numeroVersion, es_actual, ruta, hash, subido_por",
             "Historial de reemplazos"],
            ["expediente_bitacora", "id, expediente_id, tipo, usuario_id, comentario, fecha",
             "Aprobación, rechazo o corrección"],
            ["expediente_bitacora_detalle", "id, bitacora_id, documento, comentario, pagina, x/y percent",
             "Observación localizable en el PDF"],
            ["expediente_correlativo_config", "siguiente_numero, numero_inicio, digitos, anio_actual",
             "Secuencia independiente de SIAF"],
            ["producto_catalogo", "id, origen, codigo, descripcion, datos_originales (JSON)",
             "Unicidad origen+código"],
            ["producto_catalogo_config", "origen (PK), encabezados, columna_codigo, columnas_descripcion",
             "Mapeo de importación Excel"],
        ],
        "Elaboración propia a partir de backend/src/entity. Tipos físicos (varchar, timestamptz, jsonb) se omiten adrede, conforme a la consigna de no incluir scripts de tablas.",
    )


def update_indexes(doc: Document):
    """Agrega entradas a los índices estáticos y renumeran anexos 23-24 → 45-46."""
    fig_entries = [
        "Figura 23 Modelo entidad–relación de SIGEC-IGSS — Parte A: mapa de módulos",
        "Figura 24 Modelo entidad–relación — Parte B: seguridad, acceso y estructura organizacional",
        "Figura 25 Modelo entidad–relación — Parte C: módulo SIAF",
        "Figura 26 Modelo entidad–relación — Parte D: módulo de expedientes",
        "Figura 27 Modelo entidad–relación — Parte E: catálogo de productos y correlativos",
        "Figura 28 Esquema de almacenamiento híbrido de documentos en SIGEC-IGSS",
        "Figura 29 Arquitectura de la solución SIGEC-IGSS (vista de contenedores)",
        "Figura 30 Arquitectura de despliegue del piloto SIGEC-IGSS (Palín–DAF Escuintla)",
        "Figura 31 Vistas de arquitectura del SIGEC-IGSS según el modelo 4+1",
        "Figura 32 Flujo de una petición autenticada a través de las capas de SIGEC-IGSS",
        "Figura 33 Mapa de navegación del prototipo de interfaz SIGEC-IGSS",
        "Figura 34 Prototipo de interfaz — inicio de sesión",
        "Figura 35 Prototipo de interfaz — panel de administración",
        "Figura 36 Prototipo de interfaz — gestión de usuarios",
        "Figura 37 Prototipo de interfaz — panel de control del colaborador",
        "Figura 38 Prototipo de interfaz — listado de solicitudes SIAF",
        "Figura 39 Prototipo de interfaz — creación de solicitud SIAF (formato SIAF-A-01)",
        "Figura 40 Prototipo de interfaz — expedientes de compras",
        "Figura 41 Prototipo de interfaz — bandeja de revisiones DAF",
        "Figura 42 Prototipo de interfaz — análisis SIAF (cierre, trazabilidad y motivos)",
        "Figura 43 Prototipo de interfaz — actualización de códigos y productos",
        "Figura 44 Diagrama de componentes de programación de SIGEC-IGSS",
        "Figura 45 Mapa mental del proyecto SIGEC-IGSS (2026)",
        "Figura 46 Matriz operacional SIGEC-IGSS (piloto Palín–DAF Escuintla, 2026)",
    ]
    tbl_entries = [
        "Tabla 23 Dominios del modelo de datos de SIGEC-IGSS",
        "Tabla 24 Tablas principales del modelo físico y su responsabilidad",
        "Tabla 25 Relaciones y cardinalidades relevantes del diseño",
        "Tabla 26 Reglas de integridad adoptadas en el diseño de datos",
        "Tabla 27 Estructura de directorios del almacén documental",
        "Tabla 28 Políticas de archivo aplicadas en el prototipo",
        "Tabla 29 Vistas arquitectónicas y preocupaciones cubiertas",
        "Tabla 30 Especificación tecnológica por capa",
        "Tabla 31 Decisiones de arquitectura del piloto",
        "Tabla 32 Especificación de interfaces de la API por módulo",
        "Tabla 33 Principios de diseño de interfaz aplicados en el prototipo",
        "Tabla 34 Pantallas del prototipo y actor principal",
        "Tabla 35 Módulos principales de programación en el backend",
        "Tabla 36 Módulos principales de programación en el frontend",
        "Tabla 37 Trazabilidad entre módulos de programación y requerimientos",
        "Tabla 38 Diccionario de tablas del esquema sigec_igss",
    ]

    # Insert table index entries after Tabla 22
    def _clone_index_after(anchor_p, texts):
        ppr = anchor_p._p.find(qn("w:pPr"))
        for text in reversed(texts):
            el = OxmlElement("w:p")
            if ppr is not None:
                el.insert(0, deepcopy(ppr))
            anchor_p._p.addnext(el)
            para = P(el, anchor_p._parent)
            r = para.add_run(text)
            r.font.name = "Calibri"

    t22 = find_para(doc, lambda p: p.text.strip().startswith("Tabla 22 "))
    _clone_index_after(t22, tbl_entries)

    to_clear = []
    for p in doc.paragraphs:
        t = p.text.strip()
        if t.startswith("Figura 23 Mapa mental") or t.startswith("Figura 24 Matriz"):
            to_clear.append(p)
    for p in to_clear:
        parent = p._p.getparent()
        if parent is not None:
            parent.remove(p._p)

    f22 = find_para(doc, lambda p: p.text.strip().startswith("Figura 22 "))
    _clone_index_after(f22, fig_entries)

    # Renumber annex captions in body
    for p in doc.paragraphs:
        t = p.text
        if t.startswith("Figura 23") and "Mapa mental" in t:
            if p.runs:
                p.runs[0].text = p.runs[0].text.replace("Figura 23", "Figura 45", 1)
        if t.startswith("Figura 24") and "Matriz operacional" in t:
            if p.runs:
                p.runs[0].text = p.runs[0].text.replace("Figura 24", "Figura 46", 1)


def update_intro_conclusions(doc: Document):
    intro = find_para(
        doc,
        lambda p: p.text.startswith("El Capítulo II aporta el marco teórico"),
    )
    if intro.runs:
        intro.runs[0].text = (
            "El Capítulo II aporta el marco teórico (institucional, compras públicas, gestión "
            "documental, procesos, trazabilidad, indicadores, seguridad y adopción). El Capítulo III "
            "describe el marco técnico de la solución: arquitectura web, stack implementado, "
            "persistencia, despliegue, servicios auxiliares, seguridad, pruebas y continuidad "
            "operativa. El Capítulo IV especifica requerimientos, casos de uso, actividades, "
            "secuencia y reportes. El Capítulo V diseña la solución: modelo de datos por dominios, "
            "almacenamiento de documentos, arquitectura, prototipo de interfaz y módulos de "
            "programación. Las conclusiones cierran aportes y límites."
        )
        for r in intro.runs[1:]:
            r.text = ""

    sintesis = find_para(
        doc,
        lambda p: p.text.startswith("En síntesis, el SIGEC-IGSS aporta una solución factible"),
    )
    if sintesis.runs:
        sintesis.runs[0].text = (
            "En síntesis, el SIGEC-IGSS aporta una solución factible alineada con la modernización "
            "digital que el instituto ha planteado como rumbo, sin confundir el alcance del proyecto "
            "con la sustitución de plataformas estatales de contratación (p. ej. Guatecompras) ni con "
            "la resolución total de cuellos de botella que tienen también componentes organizativos y "
            "normativos ajenos al software. El valor del trabajo está en instrumentar el flujo "
            "Consultorio–DAF con roles, bitácora, versionado y métricas; en diseñar datos, "
            "arquitectura, interfaz y módulos de programación coherentes con esa instrumentación; y "
            "en demostrar —en un entorno real de gestión pública y hospitalaria-administrativa— "
            "criterios de ingeniería de software, seguridad proporcional y operación acotada al piloto "
            "definido en el documento."
        )
        for r in sintesis.runs[1:]:
            r.text = ""


def main():
    if not SRC.exists():
        raise SystemExit(f"No está el documento base: {SRC}")
    doc = Document(str(SRC))
    conclusiones = find_para(doc, lambda p: p.text.strip() == "Conclusiones")
    cur = Cursor(doc, conclusiones)
    write_chapter(cur)

    nota_matriz = None
    for p in doc.paragraphs:
        tx = p.text.strip()
        if tx.startswith("Nota:") and "matriz alinea" in tx.lower():
            nota_matriz = p
            break
    if nota_matriz is None:
        nota_matriz = find_para(doc, lambda p: p.text.strip().startswith("Anexos"))
    dummy = OxmlElement("w:p")
    nota_matriz._p.addnext(dummy)
    dummy_p = P(dummy, nota_matriz._parent)
    cur_ax = Cursor(doc, dummy_p)
    annex_dictionary(cur_ax)

    update_indexes(doc)
    update_intro_conclusions(doc)

    OUT_REPO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_REPO))
    OUT_ONEDRIVE.parent.mkdir(parents=True, exist_ok=True)
    try:
        shutil.copyfile(OUT_REPO, OUT_ONEDRIVE)
        print("Guardado:", OUT_ONEDRIVE)
    except OSError as e:
        alt = OUT_ONEDRIVE.with_name("TERCERA_ENTREGA_CAPITULO_V_COPIA.docx")
        shutil.copyfile(OUT_REPO, alt)
        print("OneDrive ocupado; copia en", alt, e)
    print("Repo:", OUT_REPO)


if __name__ == "__main__":
    main()
