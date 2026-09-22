# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\SEGUNDA ENTREGA.docx")
out = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\segunda_entrega_tail.txt")

doc = Document(str(path))
lines = []

# Dump last portion of document (from para 600)
for i, p in enumerate(doc.paragraphs):
    if i < 600:
        continue
    style = p.style.name if p.style else ""
    t = p.text
    if t.strip() or (p._p.xpath(".//*[local-name()='drawing']")):
        flag = " [IMG]" if p._p.xpath(".//*[local-name()='drawing']") else ""
        lines.append(f"{i:04d} [{style}]{flag} {t[:500]}")

# Also dump TOC area (first 140 paras)
head = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\segunda_entrega_head.txt")
hlines = []
for i, p in enumerate(doc.paragraphs):
    if i > 140:
        break
    style = p.style.name if p.style else ""
    t = p.text
    if t.strip():
        hlines.append(f"{i:04d} [{style}] {t[:300]}")
head.write_text("\n".join(hlines), encoding="utf-8")

out.write_text("\n".join(lines), encoding="utf-8")
print("tail", len(lines), "head", len(hlines))

# Heading styles used
print("--- STYLES ---")
seen = set()
for p in doc.paragraphs:
    st = p.style.name if p.style else ""
    t = p.text.strip()
    if st and ("Heading" in st or "Título" in st or t.startswith("Capítulo") or t.startswith("4.") or t.startswith("5.")):
        key = (st, t[:80])
        if key not in seen and t:
            seen.add(key)
            print(repr(st), "|", t[:90])
