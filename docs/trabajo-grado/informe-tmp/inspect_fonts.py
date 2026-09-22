# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\SEGUNDA ENTREGA.docx")
doc = Document(str(path))

# Body para samples
samples = [622, 650, 665, 835, 852, 883, 528, 621]
for i in samples:
    p = doc.paragraphs[i]
    print("="*40)
    print(i, p.style.name, repr(p.text[:80]))
    print("align", p.alignment)
    pf = p.paragraph_format
    print("indent", pf.first_line_indent, "left", pf.left_indent)
    print("line", pf.line_spacing, pf.line_spacing_rule)
    print("space", pf.space_before, pf.space_after)
    for r in p.runs[:4]:
        sz = r.font.size.pt if r.font.size else None
        print(" run font", r.font.name, sz, "bold", r.bold, "italic", r.italic, repr(r.text[:50]))

# Table formatting: shading
t = doc.tables[7]
cell = t.rows[0].cells[0]
tc = cell._tc
print("HEADER XML snippet")
print(tc.xml[:1500])

# numbered subsections in ch4
print("\n--- numbered headings ---")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t[:3] in ("4.1", "4.2", "4.3", "3.1", "3.2", "1.1", "1.2") or t.startswith("4.") or t.startswith("3."):
        if t[0].isdigit():
            print(i, p.style.name, t[:90])

print("\n--- list paragraphs ch4 ---")
for i, p in enumerate(doc.paragraphs):
    if 619 <= i <= 660 and p.style.name == "List Paragraph":
        print(i, repr(p.text[:80]))
