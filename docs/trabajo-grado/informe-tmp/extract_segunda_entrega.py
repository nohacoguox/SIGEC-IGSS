# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\SEGUNDA ENTREGA.docx")
out = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\segunda_entrega_map.txt")

doc = Document(str(path))

lines = []
lines.append(f"PARAS={len(doc.paragraphs)}")
lines.append(f"TABLES={len(doc.tables)}")
lines.append(f"SECTIONS={len(doc.sections)}")
for i, s in enumerate(doc.sections):
    lines.append(
        f"SEC{i}: page={s.page_width.cm:.2f}x{s.page_height.cm:.2f} "
        f"mL={s.left_margin.cm:.2f} mR={s.right_margin.cm:.2f} "
        f"mT={s.top_margin.cm:.2f} mB={s.bottom_margin.cm:.2f}"
    )
lines.append("")
lines.append("=== HEADINGS / KEY PARAS ===")

for i, p in enumerate(doc.paragraphs):
    style = p.style.name if p.style else ""
    t = p.text.strip()
    if not t:
        continue
    keep = False
    if style.startswith("Heading") or "Título" in style or "Heading" in style:
        keep = True
    if t.startswith(("Capítulo", "CAPÍTULO", "CAPITULO", "4.", "5.", "Figura", "Tabla", "Nota.")):
        keep = True
    if t.lower().startswith(("capítulo", "indice", "índice", "contenido", "anexo", "conclus")):
        keep = True
    if keep:
        lines.append(f"{i:04d} [{style}] {t[:240]}")

out.write_text("\n".join(lines), encoding="utf-8")
print(out)
print("written", len(lines), "lines")
