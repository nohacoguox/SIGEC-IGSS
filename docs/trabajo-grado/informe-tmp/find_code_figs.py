# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx")
doc = Document(str(path))
print("tables", len(doc.tables), "paras", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Figura 4") or t.startswith("Figura 5") or t.startswith("Extractos esenciales"):
        has_img = bool(p._p.xpath(".//*[local-name()='drawing']"))
        print(f"{i:04d} img={has_img} {t[:90]!r}")
