# -*- coding: utf-8 -*-
from docx import Document
doc = Document(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx")
for i in range(150, 175):
    t = doc.paragraphs[i].text.strip()
    if t:
        print(i, doc.paragraphs[i].style.name, repr(t[:120]))
print("--- annex captions runs ---")
for i in (1173, 1180):
    p = doc.paragraphs[i]
    print(i, [r.text for r in p.runs])
