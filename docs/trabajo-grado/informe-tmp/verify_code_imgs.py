# -*- coding: utf-8 -*-
from docx import Document
path = r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx"
doc = Document(path)
print("tables", len(doc.tables))
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() in ("Figura 45", "Figura 46", "Figura 50") or p.text.strip().startswith("Figura 45"):
        nxt = doc.paragraphs[i+2] if i+2 < len(doc.paragraphs) else None
        # find nearby img
    t = p.text.strip()
    if t.startswith("Figura 45") or t.startswith("Figura 46") or t.startswith("Figura 50"):
        if i > 200:
            print(i, repr(t[:40]))
            for j in range(i, min(i+4, len(doc.paragraphs))):
                img = bool(doc.paragraphs[j]._p.xpath(".//*[local-name()='drawing']"))
                print(" ", j, "IMG" if img else "txt", repr(doc.paragraphs[j].text[:50]))
