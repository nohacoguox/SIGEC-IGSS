# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx")
doc = Document(str(path))
print("paras", len(doc.paragraphs), "tables", len(doc.tables))

# find 5.4 and what follows
start = None
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("5.4") or "Programación" in t and t.startswith("5."):
        print(i, p.style.name, repr(t[:120]))
        start = i
    if t.startswith("Figura ") and start and i > start:
        print(i, "FIG", t[:100])
    if t.startswith("Tabla ") and start and i > start:
        print(i, "TBL", t[:100])
    if t in ("Conclusión del capítulo", "Conclusiones") and start and i > start:
        print(i, "END", t)
        break
