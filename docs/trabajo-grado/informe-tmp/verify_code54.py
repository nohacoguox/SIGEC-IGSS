# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx")
doc = Document(str(path))
out = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\verify_code54.txt")
lines = []
print("tables", len(doc.tables), "paras", len(doc.paragraphs))
start = False
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t == "Programación" or t.startswith("Programación"):
        start = True
    if start:
        if t.startswith(("Figura ", "Tabla ", "Extractos", "Criterios", "Conclus", "La programación")):
            lines.append(f"{i:04d} {t[:130]}")
        if t == "Conclusiones":
            break
# index
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith("Figura 4") and i < 180:
        lines.append("IDX " + t[:110])
    if t.startswith("Figura 5") and i < 180:
        lines.append("IDX " + t[:110])
out.write_text("\n".join(lines), encoding="utf-8")
print(out.read_text(encoding="utf-8"))
