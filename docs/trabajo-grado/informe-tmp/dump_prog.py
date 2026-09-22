# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx")
doc = Document(str(path))

out = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\sec_prog.txt")
lines = []
for i in range(1138, min(1170, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    t = p.text
    style = p.style.name if p.style else ""
    runs = []
    for r in p.runs[:4]:
        sz = r.font.size.pt if r.font.size else None
        runs.append(f"{r.font.name}/{sz}/b={r.bold}/i={r.italic}:{r.text[:40]!r}")
    flag = " [IMG]" if p._p.xpath(".//*[local-name()='drawing']") else ""
    lines.append(f"{i:04d} [{style}]{flag} {t[:400]}")
    if runs:
        lines.append("    " + " | ".join(runs))
    pf = p.paragraph_format
    lines.append(f"    align={p.alignment} first={pf.first_line_indent} line={pf.line_spacing}")

# sample caption from ch5
lines.append("\n=== SAMPLE FIG 44 ===")
for i, p in enumerate(doc.paragraphs):
    if "Figura 44" in p.text:
        lines.append(f"{i} style={p.style.name} {p.text!r}")
        for r in p.runs:
            sz = r.font.size.pt if r.font.size else None
            lines.append(f"  run {r.font.name} {sz} b={r.bold} i={r.italic} {r.text!r}")

out.write_text("\n".join(lines), encoding="utf-8")
print("wrote", len(lines))
