# -*- coding: utf-8 -*-
from pathlib import Path
import shutil
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\herramientas")
from cap5_lib import set_run

SRC = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\TERCERA_ENTREGA_CAPITULO_V.docx")
ONEDRIVE = Path(
    r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx"
)

doc = Document(str(SRC))


def flatten(p, num: int):
    title = p.text.split("\n", 1)[-1].strip()
    pPr = p._p.find(qn("w:pPr"))
    for child in list(p._p):
        if child is not pPr:
            p._p.remove(child)
    r1 = p.add_run(f"Figura {num}")
    set_run(r1, bold=True)
    r2 = p.add_run("\n" + title)
    set_run(r2, italic=True)
    print("->", num, title[:60])


for p in doc.paragraphs:
    t = p.text
    if t.startswith("Figura 23") and "Mapa mental" in t and p.style.name == "Caption":
        flatten(p, 45)
    if t.startswith("Figura 24") and "Matriz operacional" in t and p.style.name == "Caption":
        flatten(p, 46)

doc.save(str(SRC))
try:
    shutil.copyfile(SRC, ONEDRIVE)
    print("OneDrive OK")
except OSError as e:
    print("OneDrive", e)
print("done")
