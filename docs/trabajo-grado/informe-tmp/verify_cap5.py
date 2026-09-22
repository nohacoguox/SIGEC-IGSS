# -*- coding: utf-8 -*-
from pathlib import Path
from docx import Document

path = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\TERCERA_ENTREGA_CAPITULO_V.docx")
doc = Document(str(path))
print("paras", len(doc.paragraphs), "tables", len(doc.tables), "size_mb", round(path.stat().st_size/1e6, 2))

keys = []
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if not t:
        continue
    if t.startswith(("Capítulo", "5.", "Tabla ", "Figura ", "Conclus", "Anexos", "El Capítulo II")):
        keys.append(f"{i:04d} {t[:140]}")
    if t.startswith("En síntesis, el SIGEC"):
        keys.append(f"{i:04d} {t[:140]}")

out = Path(r"c:\PROYECTOS-PERSONALES\SIGEC-IGSS\docs\trabajo-grado\informe-tmp\cap5_verify.txt")
out.write_text("\n".join(keys), encoding="utf-8")
print("keys", len(keys))
print(out.read_text(encoding="utf-8")[:8000])
