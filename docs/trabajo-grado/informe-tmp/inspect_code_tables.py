# -*- coding: utf-8 -*-
from docx import Document
path = r"c:\Users\estua\OneDrive\Documents\UMG\10MO\pg2\TERCERA ENTREGA\TERCERA_ENTREGA_CAPITULO_V.docx"
doc = Document(path)
print("ntables", len(doc.tables))
for i, t in enumerate(doc.tables[-10:]):
    idx = len(doc.tables) - 10 + i
    cell = t.rows[0].cells[0].text[:80].replace("\n", " | ")
    print(idx, "cols", len(t.columns), "rows", len(t.rows), repr(cell))
